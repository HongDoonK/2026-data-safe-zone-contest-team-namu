"""R4 앱 점검 기록표(xlsx) 만들기 — D-019 §2 R4, §5 대조 동.

시트: 안내 / 동네 목록(대상 17 + 대조 10, 지도 링크, 진행 집계) / 기록(동네마다 20줄, 90_reason_table.py가 그대로 읽음) / 단가
출력: docs/report/T2_앱점검_기록표.xlsx  (수치 없음. 동 이름·이웃 점검 그룹·지도 좌표만)
"""
from pathlib import Path
from urllib.parse import quote

import geopandas as gpd
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
IN = Path(__file__).resolve().parent / "out"
BD = ROOT / "data_external" / "boundary"
DST = ROOT / "docs" / "report" / "T2_앱점검_기록표.xlsx"
N_ROWS = 20                                          # 동네마다 목록 위에서부터 20곳(규칙상 최대 30)
CONTROLS = {"달성군": ["화원읍 구라리", "구지면 창리"], "동구": ["신기동", "봉무동"], "북구": ["연경동", "사수동"],
            "서구": ["중리동", "비산동"], "수성구": ["파동", "욱수동"]}       # T2_reason_rules.md §5
GROUP = {"주변도 함께 적음": "1 먼저", "중간": "2 중간", "옆 동네가 대신할 가능성 큼": "3 나중"}
PRICE = [(g, 10000, "대구광역시청 복지과 문의(2026-10-15): 2026년 1끼 단가 1만 원, 대구 공통", "2026-10-15") for g in ["달성군", "동구", "북구", "서구", "수성구"]]

# ---- 동네 목록 -------------------------------------------------------------------------
chk = pd.read_csv(IN / "u2_denominator_check_INTERNAL.csv")
nb = pd.read_csv(IN / "neighbor_check_INTERNAL.csv").set_index("dong").verdict
tg = [d for d in chk[chk.robust.str.startswith("견고")].dong]
rows = [dict(역할="대상", key=d, 순서=GROUP.get(nb.get(d), "-")) for d in tg]
rows += [dict(역할="대조", key=f"{g}|{d}", 순서="대조") for g, ds in CONTROLS.items() for d in ds]
L = pd.DataFrame(rows)
L["구"], L["동"] = L.key.str.split("|").str[0], L.key.str.split("|").str[1]
L = L.sort_values(["구", "역할", "순서", "동"], key=lambda s: s.map({"대상": 0, "대조": 1}) if s.name == "역할" else s)

# 지도 기준점: 법정동 경계 안의 대표점(항상 경계 안) → 카카오맵 링크
br = pd.read_csv(IN / "bridge_bdong_hdong.csv", dtype=str).drop_duplicates("법정동코드")
umd = gpd.read_file(BD / "LSMD_ADM_SECT_UMD_27_202609.shp", encoding="cp949")
ri = gpd.read_file(BD / "LSMD_ADM_SECT_RI_27_202609.shp", encoding="cp949")
umd["code"], ri["code"] = umd.EMD_CD + "00", ri.RI_CD
geo = pd.concat([umd[~umd.COL_ADM_SE.isin(["27710", "27720"])][["code", "geometry"]],
                 ri[ri.COL_ADM_SE == "27710"][["code", "geometry"]]], ignore_index=True)
geo = gpd.GeoDataFrame(geo, crs=umd.crs).merge(br[["법정동코드", "bkey"]], left_on="code", right_on="법정동코드")
geo = geo.dissolve(by="bkey").reset_index()
pt = gpd.GeoSeries(geo.representative_point().values, crs=geo.crs).to_crs(4326)
ll = pd.DataFrame({"lat": pt.y.values, "lng": pt.x.values}, index=geo.bkey)
L = L.join(ll, on="key")
assert L.lat.notna().all(), L[L.lat.isna()]

# ---- 서식 ------------------------------------------------------------------------------
F = "Malgun Gothic"
INP = PatternFill("solid", fgColor="FFF2CC")          # 입력 칸
HEAD = PatternFill("solid", fgColor="1C5CAB")
CTRL = PatternFill("solid", fgColor="EEF4FC")
thin = Side(style="thin", color="D0D0D0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=1, column=i, value=c)
        x.font, x.fill, x.border = Font(name=F, bold=True, color="FFFFFF"), HEAD, BOX
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[x.column_letter].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


wb = Workbook()

# 안내
ws = wb.active
ws.title = "안내"
ws.column_dimensions["A"].width = 110
lines = [
    ("대구로 앱 점검 기록표 — 급식카드 가게가 '1끼 한도 안에서 시킬 수 있는가'", True),
    ("목적: 1순위 동네 17곳과 같은 구의 대조 동네 10곳에서, 급식카드로 주문할 수 있는 가게의 최소주문금액·배달팁·영업 여부를 같은 방식으로 적는다. 판정 규칙 D-019의 R4.", False),
    ("", False),
    ("하는 순서", True),
    ("1. '동네 목록'에서 내 이름을 '담당' 칸에 적는다. 같은 구의 대상과 대조 동네는 한 사람이 같은 날 같은 시간대에 한다.", False),
    ("2. '지도' 링크를 열어 기준점을 보고, 기준점에서 가장 가까운 아파트 단지나 주택가 주소를 대구로 앱의 배달 주소로 넣는다. 주택가가 없으면 기준점 주소. 넣은 주소를 '사용한 주소'에 적는다.", False),
    ("3. 첫 동네에서 앱이 급식카드 결제 가능 가게를 어떻게 보여 주는지(필터, 표시, 카테고리)를 찾아 '급식카드 표시 방법'에 적는다. 모두 같은 방법을 쓴다. 구분할 방법이 없으면 팀 채팅에 먼저 알린다.", False),
    ("4. 정렬은 앱 기본값 그대로 두고, 급식카드 가게 목록을 위에서부터 20곳 '기록' 시트에 적는다(20곳보다 적으면 전부). 업종은 전체.", False),
    ("5. 가게마다: 가게 이름, 업종, 최소주문금액(없으면 0), 배달팁(최소주문금액으로 주문할 때 그 주소에서 내는 금액), 영업중(지금 주문 가능 Y / 준비중·영업 종료 N).", False),
    ("6. 끝나면 '동네 목록'의 '완료'를 '예'로. 기록한 가게 수, 영업 중, 1끼 단가 이하 가게 수는 자동으로 집계된다.", False),
    ("", False),
    ("지킬 것", True),
    ("• 시간: 평일 11:00~14:00(점심). 못 하면 같은 구의 대상·대조를 같은 시간대에 하고 '시작 시각'을 적는다.", False),
    ("• 노란 칸만 입력한다. 역할·구·동·순서 칸은 바꾸지 않는다(분석 코드가 이 칸으로 묶는다).", False),
    ("• 포장만 되는 가게, 가격이 이상한 가게는 지우지 말고 '메모'에 적는다.", False),
    ("• 앱 화면은 2026-10 현재 상태라 분석 기간(2023-06~2025-06)과 다르다. 기획서에는 '현재 시점 보조 근거'로 쓴다.", False),
    ("• '단가' 시트: 2026년 1끼 단가 1만 원(대구시청 복지과 문의, 대구 공통)이 들어 있다. 앱 화면과 같은 시점의 단가다. 바뀌면 이 칸만 고친다.", False),
    ("", False),
    ("기록 시트 2행은 예시(역할 '예시')다. 분석에서 빠진다.", False),
]
for i, (t, bold) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=bold, size=13 if i == 1 else 11, color="1C5CAB" if bold else "1A1A1A")
    c.alignment = Alignment(wrap_text=True, vertical="top")

# 기록
wr = wb.create_sheet("기록")
cols = ["역할", "구", "동", "순서", "점검일", "가게", "업종", "최소주문금액", "배달팁", "영업중", "급식카드 표시", "메모"]
header(wr, cols, [7, 8, 15, 6, 11, 22, 10, 13, 10, 8, 12, 30])
ex = ["예시", "예시구", "예시동", 1, "2026-10-16", "○○분식", "분식", 12000, 3000, "Y", "Y", "포장만 되면 여기에"]
for j, v in enumerate(ex, 1):
    c = wr.cell(row=2, column=j, value=v)
    c.font, c.border = Font(name=F, italic=True, color="7F7F7F"), BOX
r = 3
for _, d in L.iterrows():
    for k in range(1, N_ROWS + 1):
        for j, v in enumerate([d.역할, d.구, d.동, k], 1):
            c = wr.cell(row=r, column=j, value=v)
            c.font, c.border = Font(name=F, color="52514E"), BOX
            if d.역할 == "대조":
                c.fill = CTRL
        for j in range(5, 13):
            c = wr.cell(row=r, column=j)
            c.fill, c.border, c.font = INP, BOX, Font(name=F)
        for j in (8, 9):
            wr.cell(row=r, column=j).number_format = "#,##0"
        r += 1
last = r - 1
dv_open = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
dv_card = DataValidation(type="list", formula1='"Y,N,모름"', allow_blank=True)
dv_num = DataValidation(type="whole", operator="between", formula1="0", formula2="100000", allow_blank=True,
                        error="원 단위 숫자만(쉼표 없이)", errorTitle="숫자")
for dv, rng in [(dv_open, f"J3:J{last}"), (dv_card, f"K3:K{last}"), (dv_num, f"H3:I{last}")]:
    wr.add_data_validation(dv)
    dv.add(rng)
wr.auto_filter.ref = f"A1:L{last}"
wr["H1"].comment = Comment("가게에 최소주문금액이 없으면 0", "T2")
wr["I1"].comment = Comment("최소주문금액으로 주문할 때 이 주소에서 내는 배달팁. 무료면 0", "T2")

# 단가
wp = wb.create_sheet("단가")
header(wp, ["구", "1끼 단가(원)", "근거", "확인일"], [10, 14, 60, 12])
for i, (g, p, src, dt) in enumerate(PRICE, 2):
    for j, v in enumerate([g, p, src, dt], 1):
        c = wp.cell(row=i, column=j, value=v)
        c.font, c.border = Font(name=F), BOX
        if j in (2, 4) or (j == 3 and p is None):
            c.fill = INP
    wp.cell(row=i, column=2).number_format = "#,##0"
pend = len(PRICE) + 1

# 동네 목록 (진행 집계는 기록·단가 시트에서 수식으로)
wl = wb.create_sheet("동네 목록", 1)
lc = ["번호", "역할", "구", "동", "순서", "지도", "사용한 주소", "담당", "점검일", "시작 시각",
      "급식카드 표시 방법", "기록한 가게", "영업 중", "1끼 단가 이하·영업 중", "완료"]
header(wl, lc, [6, 7, 8, 15, 8, 10, 30, 9, 11, 9, 22, 10, 9, 13, 7])
for i, (_, d) in enumerate(L.iterrows(), 2):
    vals = [i - 1, d.역할, d.구, d.동, d.순서]
    for j, v in enumerate(vals, 1):
        c = wl.cell(row=i, column=j, value=v)
        c.font, c.border = Font(name=F), BOX
        if d.역할 == "대조":
            c.fill = CTRL
    c = wl.cell(row=i, column=6, value="지도 열기")
    c.hyperlink = f"https://map.kakao.com/link/map/{quote(d.key.replace('|', ' '))},{d.lat:.6f},{d.lng:.6f}"
    c.font, c.border = Font(name=F, color="0563C1", underline="single"), BOX
    for j in range(7, 12):
        x = wl.cell(row=i, column=j)
        x.fill, x.border, x.font = INP, BOX, Font(name=F)
    key = f'기록!$B$3:$B${last},$C{i},기록!$C$3:$C${last},$D{i}'
    price = f'INDEX(단가!$B$2:$B${pend},MATCH($C{i},단가!$A$2:$A${pend},0))'
    fx = [f'=COUNTIFS({key},기록!$F$3:$F${last},"<>")',
          f'=COUNTIFS({key},기록!$J$3:$J${last},"Y")',
          f'=IF({price}="","단가 확인",COUNTIFS({key},기록!$J$3:$J${last},"Y",기록!$H$3:$H${last},"<="&{price}))']
    for j, f in enumerate(fx, 12):
        x = wl.cell(row=i, column=j, value=f)
        x.font, x.border, x.alignment = Font(name=F), BOX, Alignment(horizontal="center")
    x = wl.cell(row=i, column=15)
    x.fill, x.border, x.font = INP, BOX, Font(name=F)
n = len(L) + 1
dv_done = DataValidation(type="list", formula1='"예,아니오"', allow_blank=True)
wl.add_data_validation(dv_done)
dv_done.add(f"O2:O{n}")
wl.cell(row=n + 2, column=1, value="노란 칸 = 입력. 순서: 1 먼저(주변도 함께 적음) → 2 중간 → 3 나중(옆 동네가 대신할 가능성). "
                                   "대조 = 같은 구에서 평균 이상 쓰고 일반 배달 규모가 비슷한 동(D-019 §5). 지도 기준점 = 법정동 경계 안 대표점."
        ).font = Font(name=F, size=10, color="52514E")
wb.save(DST)
print("ok", DST.name, "| 동네", len(L), "| 기록 줄", last - 2)
